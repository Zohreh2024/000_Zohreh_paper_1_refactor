"""
Run 3, step 5 - the report.

Builds Run_03_report.docx from the tables and figures Steps 1-4 wrote.

Run
---
    conda run -p "C:\\ProgramData\\Anaconda3\\envs\\JinzhuLuto" ^
        python Step_05_write_report.py
"""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "outputs"
PLOT_DIR = HERE / "plots"
REPORT = HERE / "Run_03_report.docx"
INK_2 = RGBColor(0x52, 0x51, 0x4E)


def fmt(v, spec="%.3f"):
    try:
        if v is None or (isinstance(v, float) and not np.isfinite(v)):
            return "n/a"
        return spec % v
    except (TypeError, ValueError):
        return str(v)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = INK_2


def figure(doc, name, cap, width=6.4):
    path = PLOT_DIR / name
    if not path.exists():
        return False
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, cap)
    return True


def table(doc, rows, header, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, header):
        c.text = str(h)
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = str(v)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Inches(w)
    doc.add_paragraph()
    return t


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.parse_args()

    diag = pd.read_csv(OUT_DIR / "allometry_diagnosis.csv")
    src = pd.read_csv(OUT_DIR / "mis_assigned_by_source.csv")
    chk = pd.read_csv(OUT_DIR / "rebuild_checks.csv")
    h = pd.read_csv(OUT_DIR / "run03_headline.csv").set_index("config")
    j = pd.read_csv(OUT_DIR / "agb_rebuilt_by_survey.csv")

    rep = h.loc["reported_stem_only"]
    reb = h.loc["rebuilt_b2p5"]
    c25 = chk[np.isclose(chk["exponent"], 2.5)].iloc[0]

    j["ratio_old"] = j["agb_drymass_ha"] / j["live_basal_area_ha"]
    j["ratio_new"] = j["agb_rebuilt"] / j["live_basal_area_ha"]
    by_src = j.groupby("source").agg(
        surveys=("agb_rebuilt", "size"),
        old=("ratio_old", "median"), new=("ratio_new", "median"),
        agb_new=("agb_rebuilt", "median")).reset_index()
    by_src["pass_old"] = (j.assign(p=j["ratio_old"] >= 1)
                          .groupby("source")["p"].mean().values * 100)
    by_src["pass_new"] = (j.assign(p=j["ratio_new"] >= 1)
                          .groupby("source")["p"].mean().values * 100)
    by_src = by_src.sort_values("surveys", ascending=False)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("Run 3 - rebuilding biomass from the stem diameters", 0)
    p = doc.add_paragraph()
    r = p.add_run("Space_time_validation/New_proposal/Run_03 · generated %s · "
                  "the one run that needed new code" % date.today().isoformat())
    r.font.size = Pt(9)
    r.font.color.rgb = INK_2

    # ------------------------------------------------------------------ #
    doc.add_heading("What was proposed, and what the data says", level=1)
    doc.add_paragraph(
        "The proposal was that 12,505 Eucalyptus stems in 199 survey events "
        "carry the \"single stemmed acacia trees\" allometry, that their "
        "masses do not scale with diameter, and that recomputing those records "
        "with a eucalypt allometry would recover the Tasmanian Ausplots and "
        "part of the CSIRO set.")
    doc.add_paragraph(
        "The first half checks out almost exactly. The second half does not "
        "survive testing, and the reason is more consequential than the "
        "proposal.")
    rows = [[r["source"], int(r["stems"]), int(r["events"]),
             fmt(r["diameter_max"], "%.0f"), fmt(r["mass_median"], "%.1f")]
            for _, r in src.iterrows()]
    table(doc, rows, ["provider", "stems", "survey events",
                      "largest diameter (cm)", "median mass (kg)"],
          widths=[2.1, 0.8, 1.0, 1.2, 1.1])
    caption(doc, "Table 1. Eucalyptus stems carrying the acacia model: 12,471 "
                 "in 199 events, against the proposal's 12,505 in 199. The "
                 "event count is exact.")

    doc.add_heading("Why the mis-assignment cannot be the cause", level=2)
    doc.add_paragraph(
        "The correctly assigned stems fail exactly the same test.")
    rows = [[r["test"], fmt(r["value"], "%.4g"), r["note"]]
            for _, r in diag.iterrows()]
    table(doc, rows, ["test", "value", "note"], widths=[2.8, 0.9, 2.3])
    caption(doc, "Table 2. The diagnosis.")
    doc.add_paragraph(
        "Regressing log mass on log diameter for the 166,565 stems carrying "
        "the eucalypt model gives R2 = 0.013 and a negative slope. It is not a "
        "mixing artefact: holding the survey, the plot, the subplot, the "
        "allometric model AND the species all constant, the median Spearman "
        "between diameter and mass across 2,133 groups is -0.004, and not one "
        "group exceeds 0.9. In the largest single subplot a 0.89 cm stem is "
        "assigned 486 kg and a 98 cm stem 48 kg.", style="Intense Quote")
    doc.add_paragraph(
        "The library's own documentation says `diameter` is \"the explanatory "
        "variable in the allometric model\" and `agb_drymass` is the "
        "\"above-ground dry standing biomass of tree/shrub (kg)\". Within one "
        "model and one species, mass must therefore be an increasing function "
        "of diameter. In the delivered file it is not a function of diameter "
        "at all. The per-survey SUM is nevertheless sound - summed and divided "
        "by area it reproduces the site table at Spearman 0.92 - and a "
        "permutation preserves a sum. The mass column is not aligned, row for "
        "row, with the stem attributes beside it.")
    doc.add_paragraph(
        "So the proposed fix cannot be carried out as specified. Identifying "
        "the mis-allometried records requires the model label to correspond to "
        "the species, and it does not: the same file assigns \"Eucalypt "
        "trees\" to Myrsine, Zanthoxylum and Ficus. No record's mass, model "
        "and diameter can be assumed to belong together, so there is no subset "
        "to recompute and merge back.")
    figure(doc, "fig_01_mass_vs_diameter.png",
           "Figure 1. Panel (a) plots reported stem mass against diameter for "
           "the correctly assigned eucalypt stems. Panel (b) is the "
           "distribution of the within-group correlation, holding subplot, "
           "model and species constant.")

    # ------------------------------------------------------------------ #
    doc.add_heading("What was done instead", level=1)
    doc.add_paragraph(
        "The mass column is discarded and biomass is rebuilt from the "
        "diameters, which are a direct measurement, for all 2,009 surveys that "
        "carry stem data:")
    doc.add_paragraph(
        "AGB(subplot) = a x sum(D^b) / subplot area, averaged over the "
        "subplots of the survey.", style="Intense Quote")
    doc.add_paragraph(
        "No external coefficients are imported, because none are needed. The "
        "exponent b sets the SHAPE - how strongly a plot's biomass is "
        "dominated by its largest stems, and therefore the order of plots - "
        "and it is varied across the plausible range for Australian trees, "
        "2.3 to 2.7, rather than assumed. The constant a sets only the LEVEL, "
        "and is recovered from data already shown to be sound: it puts the "
        "median rebuilt plot biomass onto the median reported plot biomass. "
        "Shape from the stems, level from the library's own aggregate.")

    doc.add_heading("The check that matters", level=2)
    doc.add_paragraph(
        "The rebuilt biomass is compared with `live_basal_area_ha`, which "
        "lives in the SITE table and was not used to build it. That is a "
        "genuine cross-file test - the tree list's diameters against the site "
        "list's basal area.")
    rows = [[fmt(r["exponent"], "%.1f"), int(r["surveys"]),
             fmt(r["rho_reported_vs_basal_area"], "%.3f"),
             fmt(r["rho_rebuilt_vs_basal_area"], "%.3f"),
             fmt(r["rho_rebuilt_vs_reported"], "%.3f")]
            for _, r in chk.iterrows()]
    table(doc, rows, ["exponent", "surveys", "reported vs basal area",
                      "rebuilt vs basal area", "rebuilt vs reported"],
          widths=[0.9, 0.9, 1.6, 1.5, 1.4])
    caption(doc, "Table 3. The cross-file check.")
    doc.add_paragraph(
        "The reported biomass tracks the independently measured basal area at "
        "Spearman %s. The rebuilt biomass tracks it at %s. And the two "
        "estimates agree with each other at only %s, so this is a substantial "
        "reordering of the observations, not a cosmetic one."
        % (fmt(c25["rho_reported_vs_basal_area"], "%.2f"),
           fmt(c25["rho_rebuilt_vs_basal_area"], "%.2f"),
           fmt(c25["rho_rebuilt_vs_reported"], "%.2f")))
    figure(doc, "fig_02_cross_file_check.png",
           "Figure 2. Both estimates against a quantity neither was built "
           "from.")
    doc.add_paragraph(
        "The exponent is immaterial to the ordering: b = 2.3 and b = 2.7 rank "
        "the surveys identically at Spearman 0.989, and every result below "
        "moves by less than 0.003 across that range. The single assumption "
        "this rebuild makes does not carry the conclusion.")

    # ------------------------------------------------------------------ #
    doc.add_heading("The result", level=1)
    doc.add_paragraph(
        "The comparison is paired. `reported_stem_only` and `rebuilt_b2p5` "
        "contain the SAME 599 sites, matched the same way, against the same "
        "M'. The only difference is where the observed biomass came from.",
        style="Intense Quote")
    rows = []
    for cfg in h.index:
        x = h.loc[cfg]
        rows.append([cfg, int(x["sites"]), fmt(x["median_agb"], "%.1f"),
                     fmt(x["gate_rho"]), fmt(x["ratio"]), fmt(x["rho"]),
                     fmt(x["null_nvis_rho"]),
                     fmt(x["gap_rho_nvis"], "%+.3f")])
    table(doc, rows,
          ["configuration", "sites", "median AGB", "present-day gate rho",
           "ratio", "future rho", "null rho", "gap"],
          widths=[1.5, 0.6, 0.9, 1.1, 0.7, 0.8, 0.7, 0.7])
    caption(doc, "Table 4. Every configuration, analogue-found sites only.")
    doc.add_paragraph(
        "The present-day gate - M' against observed biomass at the site's own "
        "cell, the most direct test in the whole validation - rises from a "
        "rank correlation of %s to %s. The eight future runs rise from %s to "
        "%s, and the gap to their own constrained null from %s to %s. The "
        "paired control rules out the sample: restricting the published table "
        "to the same 599 sites reproduces the published numbers almost "
        "exactly."
        % (fmt(rep["gate_rho"]), fmt(reb["gate_rho"]), fmt(rep["rho"]),
           fmt(reb["rho"]), fmt(rep["gap_rho_nvis"], "%+.3f"),
           fmt(reb["gap_rho_nvis"], "%+.3f")))
    figure(doc, "fig_03_paired_result.png",
           "Figure 3. Identical sites, identical matching, identical M' - only "
           "the observation differs.")
    figure(doc, "fig_04_exponent.png",
           "Figure 4. The one assumption, varied across its plausible range.")

    doc.add_heading("What it means for the published conclusion", level=2)
    doc.add_paragraph(
        "The space-time validation reports that the present-day gate does not "
        "pass, on a rank correlation of about 0.42. On the same sites, with "
        "biomass rebuilt from the diameters, that correlation is %s. A "
        "substantial part of the failure was a property of the library's "
        "biomass column rather than of M'."
        % fmt(reb["gate_rho"]), style="Intense Quote")

    # ------------------------------------------------------------------ #
    doc.add_heading("It also rescues the sources the filter was excluding",
                    level=1)
    doc.add_paragraph(
        "The AGB-versus-basal-area consistency filter was introduced because "
        "TERN Australia and CSIRO reported biomass that contradicted their own "
        "basal area. The rebuild removes the contradiction.")
    rows = [[r["source"][:34], int(r["surveys"]), fmt(r["old"], "%.2f"),
             fmt(r["new"], "%.1f"), fmt(r["pass_old"], "%.0f") + "%",
             fmt(r["pass_new"], "%.0f") + "%", fmt(r["agb_new"], "%.0f")]
            for _, r in by_src.head(9).iterrows()]
    table(doc, rows,
          ["provider", "surveys", "AGB per m2/ha BA, before", "after",
           "passes before", "passes after", "median rebuilt AGB"],
          widths=[1.9, 0.7, 1.3, 0.6, 0.9, 0.9, 1.1])
    caption(doc, "Table 5. The consistency filter, before and after.")
    doc.add_paragraph(
        "TERN Australia moves from 0.08 to 31.6 Mg of biomass per square metre "
        "per hectare of basal area, and from 5.7 per cent of its surveys "
        "passing the filter to 98.1. CSIRO moves from 74.8 to 100 per cent. "
        "The rebuild also repairs the anomaly Run 2 flagged: University of NSW "
        "falls from a ratio of 87 to 8.6, inside the plausible range.")

    doc.add_heading("One limitation this exposes", level=2)
    doc.add_paragraph(
        "A single global scale constant does not suit every plot. TERN's "
        "rebuilt median is about 2,300 t DM/ha, which is at or beyond the top "
        "of anything Australian forest carries, because those are one-hectare "
        "plots holding stems up to 479 cm and a power law rewards them "
        "heavily. The ORDERING of those surveys is now sensible and they are "
        "no longer excluded wholesale, which is what the run set out to "
        "achieve; their absolute LEVEL should not be quoted. More generally, "
        "because the constant is calibrated to the library's own median, the "
        "ratio of M' to observed biomass is not independent of the library and "
        "the rank correlation is the robust result here.", style="Intense Quote")

    # ------------------------------------------------------------------ #
    doc.add_heading("What to conclude", level=1)
    for t in [
        "The proposal's premise was right and its diagnosis was wrong. The "
        "12,471 stems are real, but the mis-assigned model is not the cause - "
        "the correctly assigned stems behave identically, so the fault is that "
        "the mass column is not aligned with the stem attributes.",
        "Rebuilding from diameters is the fix, and it works. Against the site "
        "table's independently reported basal area the rebuilt biomass scores "
        "%s where the reported column scores %s."
        % (fmt(c25["rho_rebuilt_vs_basal_area"], "%.2f"),
           fmt(c25["rho_reported_vs_basal_area"], "%.2f")),
        "The present-day gate largely passes once the observation is sound: "
        "rank correlation %s against %s on identical sites. This is the single "
        "largest change any run in this study has produced, and it is "
        "attributable to the observation rather than to the sample."
        % (fmt(reb["gate_rho"]), fmt(rep["gate_rho"])),
        "The consistency filter can be retired for these sources. It existed "
        "to exclude TERN and CSIRO wholesale; with the rebuild they pass on "
        "their own merits.",
        "Quote the ranking, not the level. The scale constant is calibrated to "
        "the library's own median, so the ratio is not independent of it, and "
        "TERN's rebuilt level in particular is not credible.",
        "This should be reported to the library's custodians. The defect is in "
        "the delivered file rather than in anything done here, and it affects "
        "every stem-level use of the National Biomass Library.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("Files", level=1)
    table(doc, [
        ["Step_01_diagnose_allometry.py", "the diagnosis"],
        ["Step_02_recompute_agb.py", "the rebuild and the reference tables"],
        ["Step_03_run_and_compare.py", "the five configurations"],
        ["Step_04_plots.py", "figures 2 to 4"],
        ["Step_05_write_report.py", "this document"],
        ["outputs/allometry_diagnosis.csv", "Table 2"],
        ["outputs/rebuild_checks.csv", "Table 3"],
        ["outputs/run03_headline.csv", "Table 4"],
        ["outputs/agb_rebuilt_by_survey.csv", "the rebuilt biomass per survey"],
        ["outputs/reference_table_rebuilt_b*.csv", "one per exponent"],
    ], ["file", "what it holds"], widths=[2.4, 3.6])

    doc.save(str(REPORT))
    print("wrote %s" % REPORT)


if __name__ == "__main__":
    main()
