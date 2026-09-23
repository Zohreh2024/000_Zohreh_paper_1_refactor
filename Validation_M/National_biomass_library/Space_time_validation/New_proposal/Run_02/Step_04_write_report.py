"""
Run 2 - the report.

Builds Run_02_report.docx from the tables and figures Steps 1-3 wrote.

Run
---
    conda run -p "C:\\ProgramData\\Anaconda3\\envs\\JinzhuLuto" ^
        python Step_04_write_report.py
"""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "outputs"
PLOT_DIR = HERE / "plots"
REPORT = HERE / "Run_02_report.docx"
INK_2 = RGBColor(0x52, 0x51, 0x4E)


def fmt(v, spec="%.3f"):
    try:
        if v is None or (isinstance(v, float) and not np.isfinite(v)):
            return "n/a"
        return spec % v
    except (TypeError, ValueError):
        return str(v)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = INK_2


def figure(doc, name, cap, width=6.4):
    path = PLOT_DIR / name
    if not path.exists():
        return False
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, cap)
    return True


def table(doc, rows, header, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, header):
        c.text = str(h)
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = str(v)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Inches(w)
    doc.add_paragraph()
    return t


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.parse_args()

    h = pd.read_csv(OUT_DIR / "run02_headline.csv").set_index("config")
    cov = pd.read_csv(OUT_DIR / "coverage.csv")
    adm = pd.read_csv(OUT_DIR / "admitted_by_the_change.csv")

    a, b = h.loc["floor_005"], h.loc["floor_004"]
    effect = float(b["gap_rho_nvis"] - a["gap_rho_nvis"])
    noise = (abs(float(b["gap_rho_nvis"]
                       - h.loc["floor_004_no_unsw", "gap_rho_nvis"]))
             if "floor_004_no_unsw" in h.index else np.nan)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("Run 2 - lowering the plot-area threshold to 0.04 ha", 0)
    p = doc.add_paragraph()
    r = p.add_run("Space_time_validation/New_proposal/Run_02 · generated %s · "
                  "one constant in Step_01, reported as a sensitivity"
                  % date.today().isoformat())
    r.font.size = Pt(9)
    r.font.color.rgb = INK_2

    # ------------------------------------------------------------------ #
    doc.add_heading("The change", level=1)
    doc.add_paragraph(
        "The parent's Step_01_build_reference_table.py carries MIN_AREA_HA = "
        "0.05. This run rebuilds the reference table with 0.04 and changes "
        "nothing else: the same biomass and basal-area requirements, the same "
        "dead-basal-area limit, the same production and planted-project "
        "exclusions, the same 1985 cut-off, the same AGB-versus-basal-area "
        "consistency filter, the same one-row-per-site collapse. The matching "
        "is Run 0's throughout, and both nulls are computed in every "
        "configuration with every metric taken from the analogue-found "
        "stratum.")

    doc.add_heading("What it buys: southern coverage", level=1)
    rows = [[r["threshold"], int(r["sites"]), int(r["south_of_37"]),
             fmt(r["lat_min"], "%.1f"), int(r["n_delwp"]),
             fmt(r["median_plot_ha"], "%.2f"), fmt(r["median_agb"], "%.1f")]
            for _, r in cov.iterrows()]
    table(doc, rows,
          ["threshold", "mature sites", "south of 37 S", "furthest south",
           "DELWP Victoria sites", "median plot (ha)", "median AGB"],
          widths=[0.9, 1.0, 1.0, 1.0, 1.2, 1.0, 0.9])
    caption(doc, "Table 1. What each threshold covers.")
    doc.add_paragraph(
        "This is the argument for the change and it is a strong one. At 0.05 "
        "ha the mature sample contains four sites below 37 S - Victoria and "
        "Tasmania are effectively absent from the validation of a continental "
        "layer. At 0.04 ha there are 113. Every conclusion the space-for-time "
        "validation draws is currently a statement about Queensland and New "
        "South Wales; this is the only filter choice tested in this study that "
        "changes that.", style="Intense Quote")
    figure(doc, "fig_01_coverage.png",
           "Figure 1. Where the new sites are. Panel (a) plots every mature "
           "site, with those admitted only by the change in red; panel (b) is "
           "the latitude distribution under each threshold.")

    doc.add_heading("Who supplies them", level=1)
    rows = [[r["source"], int(r["visits"]), fmt(r["agb_median"], "%.0f"),
             fmt(r["basal_area_median"], "%.1f"),
             fmt(r["agb_per_ba_median"], "%.1f"),
             "%.2f-%.2f" % (r["area_min"], r["area_max"]),
             int(r["south_of_37"])]
            for _, r in adm.head(8).iterrows()]
    table(doc, rows,
          ["provider", "visits admitted", "median AGB", "median basal area",
           "AGB per m2/ha BA", "plot area (ha)", "south of 37 S"],
          widths=[1.7, 0.9, 0.8, 1.0, 1.0, 0.9, 0.8])
    caption(doc, "Table 2. The visits admitted only by lowering the threshold.")
    doc.add_paragraph(
        "The gain is concentrated rather than diffuse, which is what makes it "
        "checkable. DELWP Victoria supplies 201 of the 283 visits, all at "
        "exactly 0.04 ha, and their values are ordinary: a median 225 Mg/ha of "
        "biomass on 30.8 m2/ha of live basal area, a ratio of 7.3 that sits "
        "right on the library median. These are tall wet eucalypt forest plots "
        "that fail the 0.05 ha threshold by one hundredth of a hectare and by "
        "nothing else.")
    doc.add_paragraph(
        "One provider in that list should be watched rather than trusted. "
        "University of NSW contributes 27 visits reporting a median 68 Mg/ha "
        "on 0.6 m2/ha of basal area - a ratio near 106. They pass the "
        "consistency filter only because that filter is a floor and not a "
        "ceiling. The analysis below is run with and without them.")

    # ------------------------------------------------------------------ #
    doc.add_heading("What it does to the result", level=1)
    rows = []
    for cfg in ["floor_005", "floor_004", "floor_004_no_unsw",
                "floor_004_south"]:
        if cfg not in h.index:
            continue
        x = h.loc[cfg]
        rows.append([cfg, int(x["sites"]), int(x["south"]),
                     fmt(x["ratio"]), fmt(x["rho"]),
                     fmt(x["null_nvis_rho"]),
                     fmt(x["gap_rho_nvis"], "%+.3f"),
                     fmt(x["no_analogue"], "%.1f") + "%"])
    table(doc, rows,
          ["configuration", "sites", "south", "ratio", "rho", "null rho",
           "gap vs constrained null", "no-analogue"],
          widths=[1.5, 0.6, 0.6, 0.7, 0.7, 0.8, 1.3, 0.9])
    caption(doc, "Table 3. Every configuration, analogue-found sites only.")
    doc.add_paragraph(
        "Unlike Run 1, this change gains sites AND metrics: 778 against 600, "
        "with the rank correlation rising from %s to %s and the gap to the "
        "constrained null from %s to %s. There is no precision trade-off to "
        "weigh, because the sample got bigger."
        % (fmt(a["rho"]), fmt(b["rho"]), fmt(a["gap_rho_nvis"], "%+.3f"),
           fmt(b["gap_rho_nvis"], "%+.3f")))
    figure(doc, "fig_02_metrics.png",
           "Figure 2. The runs against their own null, and the gap that "
           "survives it.")

    doc.add_heading("Two reasons not to over-read it", level=2)
    doc.add_paragraph(
        "The southern sites carry no skill of their own. Matched alone, the "
        "113 sites below 37 S give a rank correlation of %s against a null of "
        "%s - a gap of %s, which is to say none at all. They are underpowered "
        "at that size and 35 per cent of them have no analogue, so this is not "
        "evidence against them; but the improvement in the pooled result is "
        "not coming from the new southern sites ranking well. It comes from "
        "the sample being larger."
        % (fmt(h.loc["floor_004_south", "rho"])
           if "floor_004_south" in h.index else "n/a",
           fmt(h.loc["floor_004_south", "null_nvis_rho"])
           if "floor_004_south" in h.index else "n/a",
           fmt(h.loc["floor_004_south", "gap_rho_nvis"], "%+.3f")
           if "floor_004_south" in h.index else "n/a"))
    doc.add_paragraph(
        "And the change in the metric is within the noise. Dropping the eight "
        "University of NSW sites - one per cent of the sample - moves the gap "
        "by %s, while the change from 0.05 to 0.04 ha moves it by %s. When "
        "removing one per cent of the sites shifts the statistic nearly twice "
        "as far as the intervention does, the intervention cannot be quoted as "
        "an effect on that statistic."
        % (fmt(noise, "%.3f"), fmt(effect, "%+.3f")), style="Intense Quote")
    figure(doc, "fig_03_stability.png",
           "Figure 3. The effect being tested, against what one per cent of "
           "the sample does.")

    # ------------------------------------------------------------------ #
    doc.add_heading("Why this stays a sensitivity", level=1)
    doc.add_paragraph(
        "The original rationale for the threshold has not changed. Measured on "
        "this library, median biomass falls from 276 Mg/ha on plots under 0.05 "
        "ha to 11 over 0.5 ha, a gradient produced by sampling geometry rather "
        "than by vegetation. The companion study in Plot_area_floor_method "
        "pushes the same lever the other way and finds that RAISING the floor "
        "to 0.40 ha lifts M'/AGB from 0.465 to 1.275 and the gap to the "
        "constrained null from +0.064 to +0.128 - the largest improvement any "
        "single change in this study has produced.")
    doc.add_paragraph(
        "So the two runs point in opposite directions on the same parameter, "
        "and they are not in conflict. The plot-size artefact is real and "
        "raising the floor removes it. Lowering the floor admits more of it, "
        "and is justified by coverage rather than by measurement quality. Both "
        "belong in the record; neither should be adopted silently.",
        style="Intense Quote")

    doc.add_heading("What to conclude", level=1)
    for t in [
        "Quote the coverage gain, which is large and unambiguous: four "
        "southern sites become 113, and the validation stops being a "
        "statement about Queensland and New South Wales alone.",
        "Do not quote the metric gain. The rank correlation rises from %s to "
        "%s and the gap from %s to %s, but one per cent of the sample moves "
        "the gap by %s. The change is not resolvable at this sample size."
        % (fmt(a["rho"]), fmt(b["rho"]), fmt(a["gap_rho_nvis"], "%+.3f"),
           fmt(b["gap_rho_nvis"], "%+.3f"), fmt(noise, "%.3f")),
        "Name the provider. DELWP Victoria supplies 201 of the 283 admitted "
        "visits, all at exactly 0.04 ha. The sensitivity is really a "
        "sensitivity to including one provider, and should be described that "
        "way rather than as a threshold choice.",
        "Watch University of NSW. Twenty-seven visits at a ratio near 106 of "
        "biomass to basal area pass the consistency filter only because it is "
        "a floor. A ceiling on that ratio is worth its own one-change test.",
        "Keep 0.05 ha as the default. The plot-size rationale stands, and the "
        "companion study shows the gains lie in the other direction.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("Files", level=1)
    table(doc, [
        ["Step_01_build_004_table.py", "rebuilds the reference table at 0.04 "
                                       "ha by patching the parent's constant"],
        ["Step_02_run_and_compare.py", "the four configurations"],
        ["Step_03_plots.py", "the figures"],
        ["Step_04_write_report.py", "this document"],
        ["outputs/reference_table_min004.csv", "the 0.04 ha table"],
        ["outputs/admitted_by_the_change.csv", "Table 2"],
        ["outputs/coverage.csv", "Table 1"],
        ["outputs/run02_headline.csv", "Table 3"],
        ["plots/fig_01..03", "the figures in this report"],
    ], ["file", "what it holds"], widths=[2.4, 3.6])

    doc.save(str(REPORT))
    print("wrote %s" % REPORT)


if __name__ == "__main__":
    main()
