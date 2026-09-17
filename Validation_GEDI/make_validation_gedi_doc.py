"""
Build the Word document explaining the Validation_GEDI folder.

Reads   outputs/source_check.csv          (written by Step_01)
        outputs/step04_New_M_2019_log.txt (fire flag shares)
        outputs/analysis/*.csv            (written by Step_05)
        common.check_alignment()          (measured live)
        CMR                               (collection ids, live, no login)
Writes  Validation_GEDI_explained.docx

Results are only summarised here; Validation_GEDI_report.docx
(make_validation_gedi_report.py) is the full write-up.
"""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

import common as C

HERE = Path(__file__).resolve().parent
OUT = HERE / "outputs"
DOCX = HERE / "Validation_GEDI_explained.docx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x60, 0x60, 0x60)
ACCENT = RGBColor(0x0B, 0x3D, 0x62)
GOOD = RGBColor(0x1E, 0x6B, 0x3A)
WARN = RGBColor(0xB4, 0x47, 0x2A)
AMBER = RGBColor(0x9A, 0x6B, 0x11)


def h(doc, t, lv):
    p = doc.add_heading(t, level=lv)
    for r in p.runs:
        r.font.color.rgb = ACCENT if lv <= 2 else INK
    return p


def para(doc, t, size=10.5, italic=False, colour=INK, after=7):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.13
    r = p.add_run(t)
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.color.rgb = colour
    return p


def rich(doc, chunks, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.13
    for t, b in chunks:
        r = p.add_run(t)
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.color.rgb = INK
    return p


def bullet(doc, t, size=10.5, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(t)
    r.font.size = Pt(size)
    r.font.color.rgb = INK
    return p


def code(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Inches(0.3)
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(8.8)
        r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def table(doc, header, rows, widths=None, size=8.8):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, n in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(n)
        r.font.bold = True
        r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def flag(doc, label, text, colour):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(label)
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = colour
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.left_indent = Inches(0.2)
    p2.paragraph_format.line_spacing = 1.13
    r2 = p2.add_run(text)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = INK
    return p2


def main():
    src = pd.read_csv(OUT / "source_check.csv")
    ana = OUT / "analysis"
    fps = pd.read_csv(ana / "footprint_summary.csv").iloc[0]
    fpr = pd.read_csv(ana / "footprints_by_region.csv").set_index("region")
    reg = pd.read_csv(ana / "by_region.csv")
    rn = reg[reg.layer == "New_M_2019"].set_index("region")
    ovl = pd.read_csv(ana / "overall_by_layer.csv")
    ov_new = ovl[(ovl.layer == "New_M_2019") & (ovl.footprints == "unburnt")].iloc[0]
    ov_all = ovl[(ovl.layer == "New_M_2019") & (ovl.footprints == "all")].iloc[0]
    log4 = (OUT / "step04_New_M_2019_log.txt").read_text(errors="ignore")
    fire_lines = [" ".join(ln.replace(":", " ").split())
                  for ln in log4.splitlines()
                  if "burnt within" in ln or "2019-20 fire" in ln]
    off = C.check_alignment(verbose=False)
    try:
        gedi = C.resolve_collection(keyword=C.GEDI_L4A_KEYWORD,
                                    prefer_version="2.1")
        v3 = C.resolve_collection(keyword=C.GEDI_L4A_KEYWORD,
                                  prefer_version="3")
        fire = C.resolve_collection(short_name=C.MODIS_BURN_SHORT)
    except Exception:
        gedi = {"short_name": "GEDI_L4A_AGB_Density_V2_1_2056",
                "id": "C2237824918-ORNL_CLOUD", "version_id": "2.1",
                "time_start": "2019-04-17", "time_end": "2025-07-09"}
        v3 = {"short_name": "GEDI_L4A_AGB_Density_V3_2508",
              "id": "C4212593885-ORNL_CLOUD", "version_id": "3"}
        fire = {"short_name": "MCD64A1", "id": "C2565786756-LPCLOUD",
                "version_id": "061", "time_start": "2000-11-01"}

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    t = doc.add_heading("Validation_GEDI: what was fixed, and how each step "
                        "works", level=0)
    for r in t.runs:
        r.font.color.rgb = ACCENT
    para(doc, "Validation_GEDI/  ·  a corrected version of the GEDI pipeline "
              "in Option_B_matched_footing/Validation/files/  ·  the original "
              "is left untouched  ·  run end to end on all six regions",
         size=9.5, italic=True, colour=MUTED)

    # ---------------- why ------------------------------------------- #
    h(doc, "Why this folder exists", 1)
    para(doc, "The GEDI plan in files/ is a good design. Its central idea — "
              "using the spread of GEDI footprints inside a single cell to "
              "estimate that cell's upper biomass envelope directly — is "
              "better than anything else available, because M is defined as a "
              "maximum and neither the National Biomass Library (one plot per "
              "location) nor the gridded GEDI L4B product (a mean) can "
              "produce such a thing.")
    para(doc, "Three defects would have prevented it from working, or from "
              "producing a trustworthy answer. This folder fixes those three "
              "and changes nothing else.")
    table(doc, ["#", "defect", "consequence if run as written", "status"],
          [["1", "Aggregation grid offset half a cell from NLUM",
            "Every cell straddled two NLUM cells; M′ and FPI were sampled "
            "on a cell boundary", "FIXED"],
           ["2", "GEDI collection short name did not resolve",
            "Zero granules returned for every region — the download would "
            "have found nothing", "FIXED"],
           ["3", "No fire masking, with Black Summer inside the record",
            "Burnt canopy could manufacture, amplify or conceal the very "
            "result being tested", "FIXED"]],
          widths=[0.3, 1.9, 2.6, 0.7])

    # ---------------- status ------------------------------------------ #
    h(doc, "Status: run end to end", 1)
    para(doc, "Steps 1 to 5 have been run on all six regions, first on "
              "vic_central as a pilot and then on the rest. GEDI L4A gave "
              "%s footprints that passed every filter, acquired %s to %s, in "
              "%s NLUM cells. Cells with at least 30 footprints and no fire "
              "in the 10 years before acquisition: %s."
         % ("{:,}".format(int(fps.footprints)), str(fps["first"])[:10],
            str(fps["last"])[:10], "{:,}".format(int(fps.cells)),
            "{:,}".format(int(ov_new.cells))))
    rows = []
    for r_ in C.REGIONS:
        rows.append([r_, "{:,}".format(int(fpr.loc[r_, "granules"])),
                     "{:,}".format(int(fpr.loc[r_, "footprints"])),
                     "{:,}".format(int(rn.loc[r_, "cells_all"])),
                     "{:,}".format(int(rn.loc[r_, "cells"])),
                     "%.1f%%" % rn.loc[r_, "exceed_p95_n30_pct"]])
    rows.append(["TOTAL", "{:,}".format(int(fps.granules_with_footprints)),
                 "{:,}".format(int(fps.footprints)),
                 "{:,}".format(int(ov_all.cells)),
                 "{:,}".format(int(ov_new.cells)),
                 "%.1f%%" % ov_new.exceed_p95_n30_pct])
    table(doc, ["region", "granules with footprints", "footprints",
                "cells ≥ 30", "cells ≥ 30, unburnt",
                "GEDI p95 > New_M_2019"], rows,
          widths=[1.4, 1.0, 1.0, 0.9, 1.0, 1.1])
    para(doc, "Granules crossing two boxes are counted in each region, so "
              "the region column does not sum to the total. The results, "
              "for both present-day M′ layers, are written up in "
              "Validation_GEDI_report.docx. In short: M′ ranks cells "
              "the right way in every region (no inversion), but it sits "
              "below GEDI's upper envelope in most closed evergreen "
              "broadleaf forest.", size=10)

    # ---------------- fix 1 ----------------------------------------- #
    doc.add_page_break()
    h(doc, "Fix 1 — the aggregation grid", 1)
    para(doc, "The original assigned each footprint to a cell with")
    code(doc, ["cell_col = floor(lon / 0.01)"])
    para(doc, "which puts cell edges at multiples of 0.01: 145.36, 145.37, "
              "145.38 and so on. NLUM's edges are at 112.925 + 0.01k, which "
              "gives 145.365, 145.375. The two grids are offset by exactly "
              "half a cell. Measured at a real longitude:")
    code(doc, ["for lon 145.3712",
               "  original 0.01-grid centre : 145.375000",
               "  NLUM-derived centre       : 145.370000",
               "  offset                    : %.6f deg = %.2f of a cell"
               % (off, off / 0.01)])
    rich(doc, [
        ("Two things went wrong as a result. Each aggregation cell straddled "
         "two NLUM cells, so the footprints pooled together did not belong to "
         "one M′ value. And the point at which M′ and FPI were sampled "
         "fell ", False),
        ("on an NLUM cell boundary", True),
        (", where the value returned depends on a rounding decision rather "
         "than on the data. This is the same trap CLAUDE.md records for "
         "ANUClimate, whose cell centres land exactly on NLUM's edges.",
         False)])
    flag(doc, "How it is fixed",
         "common.cell_index() derives the grid from NLUM's own transform, so "
         "a GEDI aggregation cell IS an NLUM cell by construction rather than "
         "by coincidence. Step 3 stamps nlum_row, nlum_col and cell_id onto "
         "every footprint at extraction time, and Step 4 samples M′ and "
         "FPI by array index rather than by coordinate. Nothing downstream "
         "can reinvent a misaligned grid, because no downstream step ever "
         "computes a cell from a longitude again.", GOOD)

    # ---------------- fix 2 ----------------------------------------- #
    h(doc, "Fix 2 — the collection short name", 1)
    para(doc, "The original pinned a short name that does not exist in NASA's "
              "metadata catalogue. A direct query with it returns zero "
              "granules for all six regions, so the download would have "
              "silently found nothing.")
    table(doc, ["", "value"],
          [["original script used",
            "GEDI_L4A_AGB_Density_V2_1   (0 granules)"],
           ["actual short name in CMR", gedi.get("short_name")],
           ["concept id", gedi.get("id")],
           ["version", gedi.get("version_id")],
           ["record", "%s to %s" % ((gedi.get("time_start") or "?")[:10],
                                    (gedi.get("time_end") or "ongoing")[:10])],
           ["a newer release also exists",
            "%s  (%s)" % (v3.get("short_name"), v3.get("id"))]],
          widths=[2.0, 4.2])
    flag(doc, "How it is fixed",
         "Rather than swap one hard-coded string for another that will also "
         "go stale, common.resolve_collection() searches CMR by keyword at "
         "run time and returns the CONCEPT ID, which is stable across "
         "releases. Step 1 prints what it resolved so a wrong match is "
         "visible immediately, and Step 3 searches by concept id rather than "
         "by name. Note that a version 3 now exists; the default remains 2.1 "
         "for reproducibility, and --version switches it.", GOOD)

    # ---------------- fix 3 ----------------------------------------- #
    doc.add_page_break()
    h(doc, "Fix 3 — fire masking", 1)
    rich(doc, [
        ("This is the defect most likely to have produced a wrong answer "
         "rather than no answer. GEDI's record begins 2019-04-18. The "
         "2019-20 fires burned roughly seven million hectares from late 2019, "
         "straight through vic_gippsland, nsw_southeast and nsw_north_seqld "
         "— three of the six target regions, all of them tall forest. ",
         False),
        ("A footprint acquired after a fire measures a burnt canopy.", True),
        (" In high-FPI tall forest that depresses biomass for reasons having "
         "nothing to do with M′.", False)])
    para(doc, "The pipeline exists to test whether M′ ranks tall forest "
              "backwards. Without a fire flag it cannot distinguish that "
              "hypothesis from “this forest burnt”, and the fires fall "
              "in exactly the places and years that matter. The effect could "
              "manufacture the inversion, amplify a real one, or hide it.")

    h(doc, "Which fire product, and why", 2)
    para(doc, "NIAFED, the obvious national choice, is not published on "
              "data.gov.au in a downloadable form — checked. MODIS MCD64A1 is, "
              "and it has a practical advantage: it is served from the same "
              "NASA Earthdata account GEDI already requires, so there is no "
              "second credential to obtain.")
    table(doc, ["property", "value"],
          [["product", "%s v%s" % (fire.get("short_name"),
                                   fire.get("version_id"))],
           ["concept id", fire.get("id")],
           ["what it gives", "per-pixel burn DATE, monthly, 500 m"],
           ["record", "%s to ongoing"
            % (fire.get("time_start") or "2000-11-01")[:10]],
           ["cost", "roughly 330 MB per region, against about 100 GB of GEDI "
                    "for the same box"]],
          widths=[1.6, 4.6])

    h(doc, "How the mask is built and used", 2)
    bullet(doc, "Step 2 downloads MCD64A1 for the chosen regions and reduces "
                "it to one raster per year on the NLUM grid: burn_year_YYYY."
                "tif, 1 where the cell burned that year. Resampling is "
                "NEAREST throughout — a burnt cell next to an unburnt one "
                "must never average into something in between.",
           numbered=True)
    bullet(doc, "It also writes last_burn_year.tif, the most recent burn year "
                "per cell, 0 for never.", numbered=True)
    bullet(doc, "Step 4 flags a footprint when its cell burned within "
                "--recovery-years of the acquisition date. The default is ten "
                "years: biomass in wet eucalypt forest does not return in "
                "one. Black Summer is flagged separately, because it is the "
                "case most likely to be questioned.", numbered=True)
    rich(doc, [
        ("Step 4 then writes ", False), ("two", True),
        (" cell files per M′ layer — gedi_cells_all_<tag>.csv and "
         "gedi_cells_unburnt_<tag>.csv — so the effect of the fires on the "
         "answer is visible rather than assumed away. If the two disagree, "
         "that is itself a result worth reporting.", False)])
    para(doc, "As run: %s. Excluding burnt footprints cuts the cells with "
              "≥ 30 footprints from %s to %s, most heavily in "
              "vic_gippsland and nsw_southeast, while the share of cells "
              "where GEDI p95 exceeds New_M_2019 moves only from %.1f%% to "
              "%.1f%%."
         % ("; ".join(fire_lines), "{:,}".format(int(ov_all.cells)),
            "{:,}".format(int(ov_new.cells)), ov_all.exceed_p95_n30_pct,
            ov_new.exceed_p95_n30_pct), size=10)
    para(doc, "MODIS MCD64A1 is HDF4, and the JinzhuLuto environment's GDAL "
              "has no HDF4 driver. Rather than modify that environment, "
              "Step 2 converts each file to a temporary GeoTIFF with the "
              "zoenv environment's gdal_translate.exe, run as a subprocess.",
         size=10)

    # ---------------- smaller fixes ---------------------------------- #
    h(doc, "Two smaller corrections", 1)
    bullet(doc, "p99 is suppressed below 100 footprints. The original allowed "
                "agbd_p99 with a minimum of 30 footprints per cell, but the "
                "99th percentile of 30 values is simply the largest of them — "
                "not a percentile, and very unstable. p95 remains the working "
                "statistic; p99 is reported only where there are enough shots "
                "to mean something.")
    bullet(doc, "The job is sized before it starts. Step 1 counts granules "
                "and estimates volume per region, because the six-region set "
                "is far larger than the original README implied.")
    rows = [[r.region, "%d" % r.gedi_granules, "%.0f GB" % r.gedi_gb_estimate,
             "%d" % r.modis_granules, "%.0f MB" % r.modis_mb_estimate]
            for _, r in src.iterrows()]
    rows.append(["TOTAL", "%d" % src.gedi_granules.sum(),
                 "%.0f GB" % src.gedi_gb_estimate.sum(),
                 "%d" % src.modis_granules.sum(),
                 "%.0f MB" % src.modis_mb_estimate.sum()])
    table(doc, ["region", "GEDI granules", "GEDI volume",
                "MODIS granules", "MODIS volume"],
          rows, widths=[1.5, 1.1, 1.0, 1.1, 1.0])
    para(doc, "These are Step 1's estimates before anything was fetched. "
              "Granules overlap between adjacent regions, so the true total "
              "is lower than the sum. Because Step 3 deletes each granule "
              "once its footprints are checkpointed, disk use stays small "
              "even though the transfer is terabyte-scale. vic_central was "
              "run first to check how many cells survive the "
              "minimum-footprint filter, and the other five followed in one "
              "resumable run.", size=10)

    # ---------------- carried over ----------------------------------- #
    h(doc, "What was deliberately left alone", 1)
    para(doc, "The original's scientific choices were sound and are carried "
              "over unchanged:")
    bullet(doc, "Power beams only. The four full-power beams penetrate dense "
                "canopy far better than the coverage beams, which matters in "
                "tall wet forest.")
    bullet(doc, "sensitivity > 0.95, stricter than the usual 0.90. A less "
                "sensitive shot may not reach the ground and will "
                "under-report biomass — precisely the error that would "
                "manufacture a false negative correlation.")
    bullet(doc, "Quality and degrade flags applied, and filtering done inside "
                "the extraction loop so only survivors are ever held in "
                "memory.")
    bullet(doc, "Granules deleted after extraction unless --keep-granules.")
    bullet(doc, "--inspect first, to confirm the HDF5 variable paths before "
                "committing to a long download.")
    bullet(doc, "The six regions, chosen to target eucalypt tall and open "
                "forest — the group where Level 2 gave an elasticity of "
                "−2.62 and a rank correlation of −0.60.")

    # ---------------- running it ------------------------------------- #
    doc.add_page_break()
    h(doc, "Running it, in order", 1)
    h(doc, "Before anything: credentials", 2)
    para(doc, "One free NASA Earthdata account covers both GEDI and the fire "
              "product. Register, then log in once and accept the ORNL DAAC "
              "and LP DAAC licence agreements — that second step is easy to "
              "skip and causes a confusing failure later.")
    para(doc, "Then store the login once with the helper. It writes "
              "N:\\Current-Users\\ZOHREH-KALAHROUDI\\.netrc (HOME is H: on "
              "this machine, so the default ~/.netrc is not used) and checks "
              "the credentials against the live service. Call python by its "
              "full path: conda run buffers input and the password prompt "
              "appears to hang.", size=10)
    code(doc, ["https://urs.earthdata.nasa.gov/users/new",
               "",
               "set PY=C:\\ProgramData\\Anaconda3\\envs\\JinzhuLuto\\python.exe",
               "%PY% setup_earthdata_login.py"])
    h(doc, "Step 1 — verify the fixes and size the job", 2)
    para(doc, "Needs no login. Prints the grid offset, the resolved "
              "collection ids, and the granule counts above.", size=10)
    code(doc, ["%PY% Step_01_check_sources.py"])
    h(doc, "Step 2 — build the fire mask", 2)
    para(doc, "All six regions by default, 2009 to 2025, so that Black "
              "Saturday (February 2009) is in the record.", size=10)
    code(doc, ["%PY% Step_02_build_fire_mask.py --check",
               "%PY% Step_02_build_fire_mask.py"])
    h(doc, "Step 3 — fetch GEDI footprints", 2)
    para(doc, "Inspect one granule first, trial a few, then run all six "
              "regions. Every granule is checkpointed, so an interrupted run "
              "resumes where it stopped.", size=10)
    code(doc, ["%PY% Step_03_fetch_gedi.py --inspect --regions vic_central",
               "%PY% Step_03_fetch_gedi.py --regions vic_central "
               "--max-granules 5",
               "%PY% Step_03_fetch_gedi.py"])
    h(doc, "Step 4 — aggregate to NLUM cells and apply the flags", 2)
    para(doc, "Once per present-day M\u2032 layer; --tag keeps the outputs "
              "apart.", size=10)
    code(doc, ["%PY% Step_04_aggregate_cells.py --tag New_M_2019",
               "%PY% Step_04_aggregate_cells.py --tag baseline_M_1985-2014 ^",
               "    --mprime ..\\Writing_paper_01\\output\\baseline_M_1985-2014.tif"])
    para(doc, "M′ defaults to Data/Processed/maxAbgM_v2/New_M_2019.tif and "
              "FPI to Random_forest_CSIRO/required_data/fpi — the same 30 "
              "annual rasters Step_08 uses as its denominator. Do not "
              "substitute FullCAM's published FPI layer; a mismatch in FPI "
              "provenance would surface as a mismatch in the biomass "
              "relationship and be misdiagnosed.", size=10)
    h(doc, "Step 5 — compare with M\u2032, then build the documents", 2)
    para(doc, "Exceedance, FPI bins, forest type, per-region tables and "
              "FPI elasticity for both layers, with five figures.", size=10)
    code(doc, ["%PY% Step_05_analyse_vs_M.py",
               "%PY% make_validation_gedi_report.py",
               "%PY% make_validation_gedi_doc.py"])

    # ---------------- not fixed -------------------------------------- #
    h(doc, "What this folder does NOT fix", 1)
    flag(doc, "Two interpretive problems remain open. They are analysis-design "
              "questions, not code defects.",
         "First, protected areas are not a random sample of tall forest. "
         "Reserves are biased toward steep, rocky, less fertile country - the "
         "land that was not worth clearing - so restricting to CAPAD "
         "categories I-IV changes the distribution of site quality, not only "
         "the disturbance history. If the rank correlation turns positive in "
         "protected forest, “clearing explains the inversion” is one "
         "reading and “protected sites span a narrower, poorer "
         "productivity range” is another. The original outcome table "
         "treats the first as established. Protection is also not the same as "
         "never-cleared: many reserves were proclaimed over previously logged "
         "forest.", AMBER)
    flag(doc, "Second, GEDI biomass is modelled, not measured.",
         "L4A derives biomass from canopy structure through allometry, so a "
         "systematic offset against field plots is expected. Characterise it "
         "where library plots and GEDI footprints coincide and report it - "
         "that comparison is a paper section in its own right, not an "
         "inconvenience. It is worth checking how well Australian eucalypt "
         "forest is represented in the L4A calibration before leaning on the "
         "absolute values.", AMBER)
    para(doc, "Neither affects whether the pipeline runs. Both affect what "
              "its output can be claimed to mean. The analysis was run "
              "without settling either. The first matters less than it did, "
              "because no inversion appeared: M′ and GEDI p95 rank cells "
              "the same way in every region, even without a protected-area "
              "restriction (Spearman %.2f to %.2f). The second is now "
              "central, because the main finding, GEDI above M′ in "
              "closed forest, could come from L4A reading high as much as "
              "from M′ being low."
         % (rn.spearman_M_vs_p95_n30.min(), rn.spearman_M_vs_p95_n30.max()))

    # ---------------- files ------------------------------------------ #
    h(doc, "Files", 1)
    table(doc, ["file", "what it does", "status"],
          [["common.py",
            "the NLUM grid, CMR collection resolution, credentials, regions",
            "imported"],
           ["setup_earthdata_login.py",
            "stores and checks the Earthdata login in .netrc", "run once"],
           ["Step_01_check_sources.py",
            "verifies both fixes, sizes the job; no login needed", "run"],
           ["Step_02_build_fire_mask.py",
            "MCD64A1 to annual burn rasters on the NLUM grid",
            "run, 6 regions"],
           ["Step_03_fetch_gedi.py",
            "GEDI L4A footprints, tagged with NLUM cell identity",
            "run, 6 regions"],
           ["Step_04_aggregate_cells.py",
            "cell summaries, M′, FPI, fire flags; writes all and unburnt",
            "run, both layers"],
           ["Step_05_analyse_vs_M.py",
            "exceedance, bins, regions, elasticity, figures", "run"],
           ["make_validation_gedi_report.py",
            "builds Validation_GEDI_report.docx (the results)", "run"],
           ["make_validation_gedi_doc.py", "builds this document", "run"],
           ["outputs/source_check.csv", "the granule and volume table", "-"],
           ["outputs/gedi_l4a_footprints.parquet",
            "every filtered footprint with its NLUM cell", "-"],
           ["outputs/gedi_cells_*.csv, analysis/, figures/",
            "Step 4 and 5 results", "-"],
           ["outputs/_vic_central_only/",
            "Step 4 and 5 results from the pilot, kept for comparison", "-"],
           ["fire/", "burn_year_YYYY.tif and last_burn_year.tif", "-"],
           ["../Option_B_matched_footing/Validation/files/",
            "the original scripts, unmodified", "-"]],
          widths=[2.2, 3.2, 1.2])

    doc.save(str(DOCX))
    print("wrote %s" % DOCX)


if __name__ == "__main__":
    main()
